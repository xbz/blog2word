import io
import unittest
from types import SimpleNamespace
from unittest.mock import patch

from docx import Document

import blog_downloader
from blog_downloader import extract_content, markdown_to_docx


class BlogDownloaderTest(unittest.TestCase):
    def test_extract_content_restores_page_h1_when_readability_omits_it(self):
        html = """
        <html><head><title>Article</title></head>
        <body><h1>Original Page H1</h1><article><p>Body text</p></article></body></html>
        """
        response = SimpleNamespace(
            text=html,
            url="https://example.com/article",
            raise_for_status=lambda: None,
        )
        readability = SimpleNamespace(summary=lambda: "<article><p>Body text</p></article>")
        with (
            patch.object(blog_downloader.requests, "get", return_value=response),
            patch.object(blog_downloader, "Document", return_value=readability),
        ):
            content, _ = extract_content(response.url)
        self.assertIn("# Original Page H1", content)

    def test_docx_preserves_heading_table_and_anchor_link(self):
        content = """# Article H1

Onderdeel | Uitleg
---|---
Link | Lees de [Thuisbatterij](https://example.com/battery)
"""
        doc = Document(io.BytesIO(markdown_to_docx(content)))

        self.assertEqual(doc.paragraphs[0].text, "Article H1")
        self.assertEqual(doc.paragraphs[0].style.name, "Heading 1")
        self.assertEqual(len(doc.tables), 1)
        self.assertEqual(doc.tables[0].cell(0, 0).text, "Onderdeel")
        self.assertEqual(doc.tables[0].cell(1, 1).text, "Lees de Thuisbatterij")

        hyperlinks = [
            relationship
            for relationship in doc.part.rels.values()
            if relationship.reltype.endswith("/hyperlink")
        ]
        self.assertEqual(len(hyperlinks), 1)
        self.assertEqual(hyperlinks[0].target_ref, "https://example.com/battery")


if __name__ == "__main__":
    unittest.main()
