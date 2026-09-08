import io
import re
import zipfile
from pathlib import Path
from urllib.parse import unquote, urljoin, urlsplit

import html2text
import requests
from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from readability import Document

URLS_FILE = "urls.txt"
OUTPUT_DIR = "output_docs"
TIMEOUT = 15
HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
}
TITLE_CLEAN_PATTERNS = [
    r"\s*-\s*[^-]+$",
    r"\s*\|\s*.+$",
    r"\s*–\s*.+$",
]


def clean_title(raw_title):
    for pattern in TITLE_CLEAN_PATTERNS:
        raw_title = re.sub(pattern, "", raw_title, flags=re.IGNORECASE)
    return raw_title.strip()


def normalize_url(url):
    """删除 URL 中横线两侧被转换器引入的空白，并编码真实空格。"""
    url = re.sub(r"\s*-\s*", "-", url.strip())
    return requests.utils.requote_uri(url)


def process_images_inline(html_content, base_url):
    soup = BeautifulSoup(html_content, "html.parser")

    headings = soup.find_all(
        ["h2", "h3", "h4"],
        string=re.compile(r"Related Blogs|You May Also Like|Recommended", re.I),
    )
    for heading in headings:
        sibling = heading.find_next_sibling()
        if sibling and (
            sibling.name in ["ul", "ol"]
            or (sibling.name == "div" and sibling.find("a"))
        ):
            sibling.decompose()
        heading.decompose()

    for div in soup.find_all("div"):
        if div.attrs is None:
            continue
        style = div.get("style", "").lower()
        if all(item in style for item in ("background-color", "border", "padding")):
            if len(div.find_all("a")) >= 3:
                div.decompose()

    for div in soup.find_all("div", class_=re.compile(r"table(-|_|)contents", re.I)):
        if div.attrs is not None:
            div.decompose()

    for img in soup.find_all("img"):
        if img.find_parent("div", class_="processed-image"):
            continue
        src = img.get("src") or img.get("data-src") or img.get("data-lazy-src") or ""
        if not src:
            img.decompose()
            continue
        alt = img.get("alt", "无").strip() or "无"
        absolute_url = normalize_url(urljoin(base_url, src))
        marker = soup.new_tag("div", **{"class": "processed-image"})
        marker.string = (
            f"@@IMG_MARKER@@\n图片\nimg_url: {absolute_url}\nalt: {alt}\n@@IMG_MARKER@@"
        )
        img.replace_with(marker)

    for div in soup.find_all("div"):
        if div.attrs is None:
            continue
        style = div.get("style", "")
        bg_match = re.search(r"background-image\s*:\s*url\(([^)]+)\)", style, re.I)
        if bg_match:
            bg_url = bg_match.group(1).strip("'\" ")
            absolute_url = normalize_url(urljoin(base_url, bg_url))
            marker = soup.new_tag("div", **{"class": "processed-image"})
            marker.string = (
                f"@@IMG_MARKER@@\n图片（背景图）\nimg_url: {absolute_url}\n"
                "alt: div背景图\n@@IMG_MARKER@@"
            )
            div.insert_before(marker)

    for link in soup.find_all("a"):
        if link.get("href"):
            link["href"] = normalize_url(urljoin(base_url, link["href"]))
    return str(soup)


def clean_markdown(text):
    text = re.sub(r"@@IMG_MARKER@@(.*?)@@IMG_MARKER@@", r"\1", text, flags=re.DOTALL)
    text = re.sub(
        r"(?m)^(img_url:\s*)(.+)$",
        lambda match: match.group(1) + normalize_url(match.group(2)),
        text,
    )
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip() + "\n"


def count_english_words(text):
    return len(re.findall(r"\b[a-zA-Z'-]+\b", text))


def extract_content(url):
    response = requests.get(url, headers=HEADERS, timeout=TIMEOUT)
    response.raise_for_status()
    canonical_url = normalize_url(response.url)
    page_soup = BeautifulSoup(response.text, "html.parser")

    raw_title = page_soup.title.get_text().strip() if page_soup.title else ""
    title = clean_title(raw_title) or "无标题"
    meta_tag = page_soup.find("meta", attrs={"name": "description"}) or page_soup.find(
        "meta", attrs={"property": "og:description"}
    )
    description = meta_tag.get("content", "").strip() if meta_tag else "无描述"
    page_h1 = page_soup.find("h1")
    page_h1_content = page_h1.get_text(" ", strip=True) if page_h1 else None

    filtered_html = process_images_inline(response.text, canonical_url)
    summary_html = Document(filtered_html).summary()
    summary_soup = BeautifulSoup(summary_html, "html.parser")
    h1_tag = summary_soup.find("h1")
    h1_content = h1_tag.get_text(" ", strip=True) if h1_tag else page_h1_content
    if h1_tag:
        h1_tag.decompose()
    processed_html = process_images_inline(str(summary_soup), canonical_url)

    converter = html2text.HTML2Text()
    converter.wrap_links = False
    converter.wrap_list_items = False
    converter.ignore_images = True
    converter.body_width = 0
    body = clean_markdown(converter.handle(processed_html))

    heading = f"# {h1_content}\n\n" if h1_content else ""
    content = (
        f"**title**: {title}\n\n"
        f"**description**: {description}\n\n"
        f"URL：{canonical_url}\n\n"
        f"{heading}{body}"
    )
    clean_text = re.sub(r"(?m)^(图片.*|img_url:.*|alt:.*)$", "", content)
    return content, count_english_words(clean_text)


def filename_from_url(url):
    slug = unquote(urlsplit(url).path.rstrip("/").split("/")[-1])
    return re.sub(r"[^\w.-]+", "_", slug).strip("_.") or "untitled"


def strip_markdown_inline(text):
    text = re.sub(r"!\[([^\]]*)\]\([^)]+\)", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    text = re.sub(r"(\*\*|__|\*|_|`|~~)", "", text)
    return text


def add_hyperlink(paragraph, text, url):
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "3157D5")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend([color, underline])
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.extend([properties, text_element])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_markdown_inline(paragraph, text):
    """将链接和基础强调标记渲染为 Word runs，不保留 Markdown 符号。"""
    pattern = re.compile(r"\[([^\]]+)\]\((https?://[^)]+)\)|(\*\*|__)(.+?)\3|(\*|_)(.+?)\5")
    position = 0
    for match in pattern.finditer(text):
        if match.start() > position:
            paragraph.add_run(strip_markdown_inline(text[position : match.start()]))
        if match.group(1):
            add_hyperlink(paragraph, strip_markdown_inline(match.group(1)), match.group(2))
        elif match.group(4):
            paragraph.add_run(strip_markdown_inline(match.group(4))).bold = True
        else:
            paragraph.add_run(strip_markdown_inline(match.group(6))).italic = True
        position = match.end()
    if position < len(text):
        paragraph.add_run(strip_markdown_inline(text[position:]))


def split_markdown_table_row(line):
    line = line.strip().strip("|")
    return [cell.strip().replace(r"\|", "|") for cell in re.split(r"(?<!\\)\|", line)]


def is_markdown_table_separator(line):
    cells = split_markdown_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells)


def add_markdown_table(doc, lines):
    rows = [split_markdown_table_row(line) for line in lines if not is_markdown_table_separator(line)]
    column_count = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=column_count)
    table.style = "Table Grid"
    for row_index, row in enumerate(rows):
        for column_index, text in enumerate(row):
            paragraph = table.cell(row_index, column_index).paragraphs[0]
            add_markdown_inline(paragraph, text)
            if row_index == 0:
                for run in paragraph.runs:
                    run.bold = True


def markdown_to_docx(content):
    """生成带标题、表格和可点击锚文字链接的 Word 文件。"""
    doc = DocxDocument()
    lines = content.splitlines()
    index = 0
    while index < len(lines):
        raw_line = lines[index]
        line = raw_line.strip()
        if not line:
            index += 1
            continue
        if (
            "|" in line
            and index + 1 < len(lines)
            and is_markdown_table_separator(lines[index + 1])
        ):
            table_lines = [line, lines[index + 1]]
            index += 2
            while index < len(lines) and "|" in lines[index] and lines[index].strip():
                table_lines.append(lines[index])
                index += 1
            add_markdown_table(doc, table_lines)
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", line)
        if heading:
            paragraph = doc.add_heading(level=len(heading.group(1)))
            add_markdown_inline(paragraph, heading.group(2))
        elif re.match(r"^[-*+]\s+", line):
            paragraph = doc.add_paragraph(style="List Bullet")
            add_markdown_inline(paragraph, re.sub(r"^[-*+]\s+", "", line))
        elif re.match(r"^\d+\.\s+", line):
            paragraph = doc.add_paragraph(style="List Number")
            add_markdown_inline(paragraph, re.sub(r"^\d+\.\s+", "", line))
        else:
            paragraph = doc.add_paragraph()
            add_markdown_inline(paragraph, line)
        index += 1
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def build_zip(urls):
    output = io.BytesIO()
    results = []
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        for index, url in enumerate(urls, 1):
            content, word_count = extract_content(url)
            base = filename_from_url(url)
            archive.writestr(f"{base}.md", content)
            archive.writestr(f"{index}. {base}.docx", markdown_to_docx(content))
            results.append((url, word_count))
    return output.getvalue(), results


def run_cli():
    urls_path = Path(URLS_FILE)
    if not urls_path.exists():
        raise SystemExit(f"错误：找不到URL列表文件 {URLS_FILE}")
    urls = [line.strip() for line in urls_path.read_text(encoding="utf-8").splitlines() if line.strip()]
    output_dir = Path(OUTPUT_DIR)
    output_dir.mkdir(parents=True, exist_ok=True)
    total_words = 0
    print(f"正在处理 {len(urls)} 个URL...")
    for index, url in enumerate(urls, 1):
        try:
            content, word_count = extract_content(url)
            base = filename_from_url(url)
            (output_dir / f"{base}.md").write_text(content, encoding="utf-8")
            (output_dir / f"{index}. {base}.docx").write_bytes(markdown_to_docx(content))
            total_words += word_count
            print(f"已处理：{base.ljust(40)} ({word_count} 单词)")
        except Exception as exc:
            print(f"抓取失败 {url}: {exc}")
    print(f"\n处理完成！结果保存至 [{OUTPUT_DIR}] 目录")
    print(f"总单词数：{total_words:,}")
